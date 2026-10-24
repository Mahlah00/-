# -*- coding: utf-8 -*-
"""
سلف تلگرام (Userbot) با Telethon

قابلیت‌ها:
 1) ماشین‌حساب            .calc 2+2*3   |  .حساب (5+3)^2
 2) ذخیره مدیای نابودشونده (خودکار + دستی با .save روی ریپلای)
 3) لاگ همه‌ی چت‌های پیوی در کانال (با ادیت و حذف)
 4) بکاپ هر ۲۴ ساعته‌ی گروه‌های انتخابی به صورت HTML / Excel

دستورات (فقط از اکانت خودت کار می‌کنند):
 .calc <عبارت>
 .save                       (ریپلای روی مدیای نابودشونده)
 .backup add [html|xlsx|both] [آیدی/یوزرنیم]
 .backup remove [آیدی/یوزرنیم]
 .backup list
 .backup now [آیدی/یوزرنیم]
 .id                         (آیدی چت فعلی را در Saved Messages می‌فرستد)
 .help
"""
import os
import re
import ast
import math
import time
import html
import sqlite3
import asyncio
import logging
import operator
import functools
import secrets
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from telethon import TelegramClient, events, utils
from telethon.tl.types import User
import aiohttp
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

try:
    import jdatetime
except ImportError:  # اگر نصب نبود تاریخ میلادی نشان می‌دهد
    jdatetime = None

# ============================ تنظیمات ============================
API_ID = 774143                      # از my.telegram.org
API_HASH = "62a87b8c12cf0c91f8d5312bfac42371"  # از my.telegram.org
SESSION = "mohammad"

# فقط مقدار پیش‌فرض ("me" = Saved Messages).
# مقصد لاگ و بکاپ را با دستور .setlog و .setbackup تنظیم کن، نیازی به ثبت در سورس نیست.
LOG_CHANNEL = "me"
BACKUP_TARGET = LOG_CHANNEL          # مقصد فایل‌های بکاپ
TTL_TARGET = "me"                    # مقصد مدیای نابودشونده

PREFIX = "."
LOG_OUTGOING = True      # پیام‌هایی که خودت در پیوی می‌فرستی هم لاگ شود
IGNORE_BOTS = True       # پیوی ربات‌ها لاگ نشود
CACHE_MEDIA = True       # مدیای پیوی‌ها ذخیره شود (برای بازیابی بعد از حذف)
MAX_MEDIA_MB = 20        # فایل‌های بزرگ‌تر کش نمی‌شوند
RETENTION_DAYS = 30      # پاک‌سازی کش و دیتابیس پیام‌های قدیمی‌تر
TIMEZONE = "Asia/Tehran"
CACHE_DIR = "cache"
TMP_DIR = "tmp"
DB_FILE = "self.db"
# ================================================================

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("self")

os.makedirs(CACHE_DIR, exist_ok=True)
os.makedirs(TMP_DIR, exist_ok=True)

TZ = ZoneInfo(TIMEZONE)
P = re.escape(PREFIX)
COMMANDS = {"calc", "c", "حساب", "save", "backup", "id", "help",
            "setlog", "setbackup", "settings", "pay", "bale",
            "ai", "claude", "gpt", "grok", "solve", "حل", "tr", "aiset", "config"}

client = TelegramClient(SESSION, API_ID, API_HASH)
client.parse_mode = None  # متن‌ها بدون مارک‌داون ارسال شوند

ME = None
LOG_ENT = None
BACKUP_ENT = None
SELF_EDITED = set()  # آیدی پیام‌هایی که خود سلف ادیت کرده (لاگ نشوند)

# ----------------------------- دیتابیس -----------------------------
db = sqlite3.connect(DB_FILE)
db.row_factory = sqlite3.Row
db.executescript("""
CREATE TABLE IF NOT EXISTS messages(
  msg_id INTEGER PRIMARY KEY, chat_id INTEGER, sender_id INTEGER,
  sender_name TEXT, username TEXT, partner_name TEXT, partner_id INTEGER,
  out INTEGER, text TEXT, media_path TEXT, ts INTEGER, ttl INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS backups(
  chat_id INTEGER PRIMARY KEY, title TEXT, fmt TEXT, last_backup INTEGER);
""")
# مهاجرت دیتابیس قدیمی: ستون‌های مربوط به پیام لاگ داخل کانال
for _col, _typ in (("log_msg_id", "INTEGER"), ("log_text", "TEXT")):
    try:
        db.execute(f"ALTER TABLE messages ADD COLUMN {_col} {_typ}")
    except sqlite3.OperationalError:
        pass
db.execute("CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT)")
db.commit()


# ------------------------------ ابزارها ------------------------------
def fmt_date(ts):
    dt = datetime.fromtimestamp(ts, TZ)
    if jdatetime:
        return jdatetime.datetime.fromgregorian(datetime=dt.replace(tzinfo=None)).strftime("%Y/%m/%d")
    return dt.strftime("%Y/%m/%d")


def fmt_time(ts):
    return datetime.fromtimestamp(ts, TZ).strftime("%H:%M:%S")


def who(entity):
    if entity is None:
        return "نامشخص", "—"
    name = utils.get_display_name(entity) or "بدون‌نام"
    uname = getattr(entity, "username", None)
    return name, (f"@{uname}" if uname else "—")


def info_block(sname, uname, sid, pname, pid, out, ts):
    lines = [
        f"👤 ارسال‌کننده: {sname}",
        f"🆔 آیدی: {sid} | {uname}",
    ]
    if out:
        lines.append(f"💬 گفتگو با: {pname} ({pid})")
    lines += [f"🕒 زمان ارسال: {fmt_time(ts)}", f"📅 تاریخ ارسال: {fmt_date(ts)}"]
    return "\n".join(lines)


def media_kind(m):
    if m.photo: return "عکس"
    if m.sticker: return "استیکر"
    if m.voice: return "ویس"
    if m.video_note: return "ویدیو مسیج"
    if m.video: return "ویدیو"
    if m.gif: return "گیف"
    if m.audio: return "آهنگ"
    if m.document: return "فایل"
    if m.poll: return "نظرسنجی"
    if m.geo: return "لوکیشن"
    if m.contact: return "مخاطب"
    return ""


def is_command(text):
    if not text.startswith(PREFIX):
        return False
    first = text[len(PREFIX):].split(None, 1)
    return bool(first) and first[0].lower() in COMMANDS


LOG_LOCK = asyncio.Lock()


def ordered(fn):
    """هندلرها را پشت‌سرهم اجرا می‌کند تا ترتیب پیام‌ها در کانال
    دقیقاً مثل ترتیب چت باشد (اول پیام طرف مقابل، بعد پیام شما)."""
    @functools.wraps(fn)
    async def wrapper(event):
        async with LOG_LOCK:
            return await fn(event)
    return wrapper


def set_log(msg_id, log_id, log_text):
    db.execute("UPDATE messages SET log_msg_id=?, log_text=? WHERE msg_id=?",
               (log_id, log_text, msg_id))
    db.commit()


async def send_log(text, file=None, reply_to=None):
    """پیام را در کانال می‌فرستد و آیدی اولین پیام ارسال‌شده را برمی‌گرداند."""
    first = None
    try:
        if file and os.path.exists(file):
            m = await client.send_file(LOG_ENT, file, caption=text[:1000], reply_to=reply_to)
            first = m
            rest = text[1000:]
        else:
            rest = text
        for i in range(0, len(rest), 4000):
            m = await client.send_message(LOG_ENT, rest[i:i + 4000], reply_to=reply_to)
            if first is None:
                first = m
    except Exception:
        log.exception("send_log failed")
    return first.id if first is not None else None


async def to_me(event, text):
    """پیام دستور را پاک می‌کند و جواب را در Saved Messages می‌فرستد."""
    try:
        await event.delete()
    except Exception:
        pass
    await client.send_message("me", text)


async def cache_media(msg):
    if not CACHE_MEDIA or not msg.file:
        return None
    size = getattr(msg.file, "size", None)
    if size and size > MAX_MEDIA_MB * 1024 * 1024:
        return None
    try:
        return await msg.download_media(file=os.path.join(CACHE_DIR, f"{msg.id}"))
    except Exception:
        log.exception("cache_media failed")
        return None


def store(msg_id, chat_id, sender, partner, out, text, path, ts, ttl=0):
    sname, uname = who(sender)
    pname, _ = who(partner)
    db.execute(
        """INSERT INTO messages(msg_id,chat_id,sender_id,sender_name,username,
        partner_name,partner_id,out,text,media_path,ts,ttl)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(msg_id) DO UPDATE SET chat_id=excluded.chat_id,
        sender_id=excluded.sender_id, sender_name=excluded.sender_name,
        username=excluded.username, partner_name=excluded.partner_name,
        partner_id=excluded.partner_id, out=excluded.out, text=excluded.text,
        media_path=excluded.media_path, ts=excluded.ts, ttl=excluded.ttl""",
        (msg_id, chat_id, getattr(sender, "id", 0), sname, uname, pname,
         getattr(partner, "id", 0), int(out), text, path, ts, ttl),
    )
    db.commit()


# ------------------------- ۱) ماشین‌حساب -------------------------
DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "0123456789" * 2)


def _pow(a, b):
    if abs(b) > 1000:
        raise ValueError("توان خیلی بزرگ است")
    if a and abs(a) > 1 and abs(b) * math.log10(abs(a)) > 1000:
        raise ValueError("نتیجه خیلی بزرگ است")
    return a ** b


def _fact(n):
    if n < 0 or n > 500 or int(n) != n:
        raise ValueError("فاکتوریل نامعتبر")
    return math.factorial(int(n))


BIN = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
       ast.Div: operator.truediv, ast.FloorDiv: operator.floordiv,
       ast.Mod: operator.mod, ast.Pow: _pow}
UN = {ast.USub: operator.neg, ast.UAdd: operator.pos}
FUNCS = {"sqrt": math.sqrt, "sin": math.sin, "cos": math.cos, "tan": math.tan,
         "log": math.log, "log10": math.log10, "abs": abs, "round": round,
         "floor": math.floor, "ceil": math.ceil, "fact": _fact}
CONSTS = {"pi": math.pi, "e": math.e}


def _eval(node):
    if isinstance(node, ast.Expression):
        return _eval(node.body)
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)) \
            and not isinstance(node.value, bool):
        return node.value
    if isinstance(node, ast.BinOp) and type(node.op) in BIN:
        return BIN[type(node.op)](_eval(node.left), _eval(node.right))
    if isinstance(node, ast.UnaryOp) and type(node.op) in UN:
        return UN[type(node.op)](_eval(node.operand))
    if isinstance(node, ast.Name) and node.id in CONSTS:
        return CONSTS[node.id]
    if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) \
            and node.func.id in FUNCS and not node.keywords:
        return FUNCS[node.func.id](*[_eval(a) for a in node.args])
    raise ValueError("عبارت نامعتبر")


def calculate(expr):
    e = (expr.translate(DIGITS).replace("×", "*").replace("÷", "/")
         .replace("٫", ".").replace("^", "**").strip())
    if len(e) > 200:
        raise ValueError("عبارت خیلی طولانی است")
    r = _eval(ast.parse(e, mode="eval"))
    if isinstance(r, float):
        if r.is_integer() and abs(r) < 1e15:
            return str(int(r))
        return format(r, ".12g")
    return str(r)


@client.on(events.NewMessage(outgoing=True, pattern=rf"(?s)^{P}(?:calc|c|حساب)\s+(.+)$"))
async def cmd_calc(event):
    expr = event.pattern_match.group(1).strip()
    try:
        res = calculate(expr)
        out = f"🧮 {expr}\n= {res}"
    except ZeroDivisionError:
        out = f"🧮 {expr}\n❌ تقسیم بر صفر"
    except Exception as ex:
        out = f"🧮 {expr}\n❌ خطا: {ex}"
    SELF_EDITED.add(event.id)
    await event.edit(out)


# ------------------- ۲) ذخیره‌ی مدیای نابودشونده -------------------
async def save_ttl(msg, sender, partner):
    path = None
    try:
        path = await msg.download_media(file=os.path.join(TMP_DIR, f"ttl_{msg.id}"))
        if not path:
            return False
        sname, uname = who(sender)
        ts = int(msg.date.timestamp())
        cap = ("💣 مدیای نابودشونده ذخیره شد\n"
               + info_block(sname, uname, getattr(sender, "id", "؟"), sname,
                            getattr(partner, "id", "؟"), False, ts))
        await client.send_file(TTL_TARGET, path, caption=cap[:1000],
                               voice_note=bool(msg.voice))
        os.remove(path)
        return True
    except Exception:
        log.exception("save_ttl failed")
        return False


@client.on(events.NewMessage(outgoing=True, pattern=rf"^{P}save$"))
async def cmd_save(event):
    r = await event.get_reply_message()
    if not r or not r.media:
        return await to_me(event, "⚠️ دستور .save را روی یک پیام مدیادار ریپلای کن.")
    sender = await r.get_sender()
    partner = await event.get_chat()
    ok = await save_ttl(r, sender, partner)
    await to_me(event, "✅ ذخیره شد." if ok else "❌ ذخیره نشد (شاید مدیا دیگر در دسترس نیست).")


# --------------------- ۳) لاگ پیوی‌ها در کانال ---------------------
@client.on(events.NewMessage())
@ordered
async def on_new(event):
    if not event.is_private or event.chat_id == ME.id:
        return
    msg = event.message
    text = msg.raw_text or ""
    out = bool(event.out)
    if out and (not opt("outgoing") or is_command(text)):
        return
    sender = await event.get_sender()
    if not opt("bots") and getattr(sender, "bot", False):
        return
    partner = await event.get_chat()
    ts = int(msg.date.timestamp())
    sname, uname = who(sender)
    pname, _ = who(partner)

    ttl = getattr(msg.media, "ttl_seconds", None) if msg.media else None
    path = None
    kind = media_kind(msg)
    if ttl:
        saved = False if (out or not opt("ttl")) else await save_ttl(msg, sender, partner)
        kind = f"💣 مدیای نابودشونده ({kind or 'مدیا'})" + (" — ذخیره شد" if saved else "")
    elif msg.media and opt("media") and (out or opt("incoming")):
        path = await cache_media(msg)
    if not out and not opt("incoming"):
        return  # لاگ دریافتی‌ها خاموش است (فقط مدیای نابودشونده در صورت روشن بودن ذخیره شد)

    body = text
    if kind:
        body = f"[{kind}]\n{text}".strip()
    store(msg.id, event.chat_id, sender, partner, out, text, path, ts, 1 if ttl else 0)

    head = "📤 پیام ارسالی شما" if out else "📩 پیام جدید"
    info = info_block(sname, uname, getattr(sender, "id", "؟"), pname,
                      getattr(partner, "id", "؟"), out, ts)
    log_text = f"{head}\n{info}\n\n{body}"
    log_id = await send_log(log_text, path)
    set_log(msg.id, log_id, log_text)


@client.on(events.MessageEdited())
@ordered
async def on_edit(event):
    if not event.is_private or event.chat_id == ME.id:
        return
    msg = event.message
    if msg.id in SELF_EDITED:
        SELF_EDITED.discard(msg.id)
        return
    out = bool(event.out)
    if not opt("edits") or (out and not opt("outgoing")) or (not out and not opt("incoming")):
        return
    new = msg.raw_text or ""
    row = db.execute("SELECT * FROM messages WHERE msg_id=?", (msg.id,)).fetchone()
    old = row["text"] if row else None
    if old is not None and old == new:
        return
    sender = await event.get_sender()
    if not opt("bots") and getattr(sender, "bot", False):
        return
    partner = await event.get_chat()
    ts = int(msg.date.timestamp())
    sname, uname = who(sender)
    pname, _ = who(partner)
    ets = int(msg.edit_date.timestamp()) if msg.edit_date else int(time.time())

    log_id = row["log_msg_id"] if row else None
    log_text = row["log_text"] if row else None
    store(msg.id, event.chat_id, sender, partner, out, new,
          row["media_path"] if row else None, ts, row["ttl"] if row else 0)

    # ادیت روی همان پیام داخل کانال:
    #   متن قبلی
    #   -------------------
    #   ادیت شده
    #
    #   متن جدید
    if log_id and log_text:
        new_log = (f"{log_text}\n-------------------\n"
                   f"✏️ ادیت شده ({fmt_time(ets)})\n\n{new}")
        try:
            await client.edit_message(LOG_ENT, log_id, new_log)
            set_log(msg.id, log_id, new_log)
            return
        except Exception:
            # مثلاً متن بلندتر از حد مجاز تلگرام شده؛ زیر همان پیام ریپلای می‌شود
            log.warning("in-place edit failed, sending reply instead", exc_info=True)

    info = info_block(sname, uname, getattr(sender, "id", "؟"), pname,
                      getattr(partner, "id", "؟"), out, ts)
    old_txt = old if old is not None else "❔ نامشخص (قبل از اجرای سلف ثبت نشده)"
    text = (f"✏️ ویرایش شده\n{info}\n"
            f"🛠 زمان ویرایش: {fmt_time(ets)} - {fmt_date(ets)}\n\n"
            f"📝 پیام قبل:\n{old_txt}\n\n📝 پیام بعد:\n{new}")
    new_id = await send_log(text, reply_to=log_id)
    if not log_id:
        set_log(msg.id, new_id, text)


@client.on(events.MessageDeleted())
@ordered
async def on_delete(event):
    if event.chat_id is not None or not opt("deletes"):  # فقط پیوی (در پیوی chat_id خالی است)
        return
    for mid in event.deleted_ids:
        row = db.execute("SELECT * FROM messages WHERE msg_id=?", (mid,)).fetchone()
        if not row:
            continue
        db.execute("DELETE FROM messages WHERE msg_id=?", (mid,))
        db.commit()
        if row["ttl"]:  # مدیای نابودشونده طبیعتاً حذف می‌شود؛ لاگ نمی‌خواهد
            continue
        info = info_block(row["sender_name"], row["username"], row["sender_id"],
                          row["partner_name"], row["partner_id"], bool(row["out"]), row["ts"])
        now = int(time.time())
        text = (f"🗑 پیام\n{info}\n\n{row['text'] or ''}\n\n"
                f"❌ پاک شده (در {fmt_time(now)} - {fmt_date(now)})")
        await send_log(text, row["media_path"])
        p = row["media_path"]
        if p and os.path.exists(p):
            try:
                os.remove(p)
            except OSError:
                pass


def purge_old():
    cutoff = time.time() - RETENTION_DAYS * 86400
    for r in db.execute("SELECT media_path FROM messages WHERE ts < ?", (cutoff,)).fetchall():
        p = r["media_path"]
        if p and os.path.exists(p):
            try:
                os.remove(p)
            except OSError:
                pass
    db.execute("DELETE FROM messages WHERE ts < ?", (cutoff,))
    db.commit()


# ------------------------- ۴) بکاپ گروه‌ها -------------------------
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(s):
    return ILLEGAL.sub("", s or "")[:32000]


def build_html(title, rows, since, until):
    nl = "\n"
    trs = []
    for i, r in enumerate(rows, 1):
        txt = html.escape(r["text"]).replace(nl, "<br>")
        trs.append(
            f"<tr><td>{i}</td><td>{r['date']}</td><td>{r['time']}</td>"
            f"<td>{html.escape(r['name'])}</td><td dir='ltr'>{r['sid']}</td>"
            f"<td>{txt}</td><td>{r['kind']}</td><td>{r['edited']}</td></tr>"
        )
    return f"""<!DOCTYPE html>
<html lang="fa" dir="rtl"><head><meta charset="utf-8">
<title>{html.escape(title)}</title>
<style>
body{{font-family:Tahoma,Vazirmatn,sans-serif;background:#f4f6f8;margin:20px;color:#222}}
h2{{margin-bottom:4px}} .sub{{color:#666;margin-bottom:16px}}
table{{border-collapse:collapse;width:100%;background:#fff}}
th,td{{border:1px solid #ddd;padding:8px;vertical-align:top;font-size:14px}}
th{{background:#2a7de1;color:#fff}} tr:nth-child(even){{background:#f9fbfd}}
</style></head><body>
<h2>بکاپ چت: {html.escape(title)}</h2>
<div class="sub">از {fmt_date(since)} {fmt_time(since)} تا {fmt_date(until)} {fmt_time(until)} — {len(rows)} پیام</div>
<table><tr><th>#</th><th>تاریخ</th><th>زمان</th><th>فرستنده</th><th>آیدی</th><th>متن</th><th>نوع</th><th>وضعیت</th></tr>
{nl.join(trs)}
</table></body></html>"""


def build_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Backup"
    ws.sheet_view.rightToLeft = True
    ws.append(["#", "تاریخ", "زمان", "فرستنده", "آیدی", "متن", "نوع", "وضعیت"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, r in enumerate(rows, 1):
        ws.append([i, r["date"], r["time"], clean(r["name"]), r["sid"],
                   clean(r["text"]), r["kind"], r["edited"]])
    for col, w in zip("ABCDEFGH", (6, 12, 10, 22, 14, 70, 12, 12)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


async def run_backup(chat_id, title, fmt, since_ts, until_ts):
    since = datetime.fromtimestamp(since_ts, timezone.utc)
    until = datetime.fromtimestamp(until_ts, timezone.utc)
    rows, cache = [], {}
    async for m in client.iter_messages(chat_id, reverse=True, offset_date=since):
        if m.date > until:
            break
        if m.action:  # پیام‌های سرویسی (جوین/لفت و...)
            continue
        sid = m.sender_id
        if sid not in cache:
            try:
                s = await m.get_sender()
            except Exception:
                s = None
            cache[sid] = utils.get_display_name(s) if s else title
        ts = int(m.date.timestamp())
        rows.append({
            "date": fmt_date(ts), "time": fmt_time(ts), "name": cache[sid],
            "sid": sid or "", "text": m.raw_text or "", "kind": media_kind(m),
            "edited": "ویرایش‌شده" if m.edit_date else "",
        })

    cap = (f"📦 بکاپ چت «{title}»\n"
           f"📅 {fmt_date(since_ts)} {fmt_time(since_ts)} تا {fmt_date(until_ts)} {fmt_time(until_ts)}\n"
           f"💬 {len(rows)} پیام")
    if not rows:
        await client.send_message(BACKUP_ENT, cap + "\n(پیامی نبود)")
        return 0

    safe = re.sub(r"[^\w\-]+", "_", title)[:30] or "chat"
    base = os.path.join(TMP_DIR, f"backup_{safe}_{fmt_date(until_ts).replace('/', '-')}")
    files = []
    if fmt in ("html", "both"):
        with open(base + ".html", "w", encoding="utf-8") as f:
            f.write(build_html(title, rows, since_ts, until_ts))
        files.append(base + ".html")
    if fmt in ("xlsx", "both"):
        build_xlsx(base + ".xlsx", rows)
        files.append(base + ".xlsx")
    for fp in files:
        await client.send_file(BACKUP_ENT, fp, caption=cap, force_document=True)
        os.remove(fp)
    return len(rows)


async def scheduler():
    while True:
        try:
            now = int(time.time())
            for r in db.execute("SELECT * FROM backups").fetchall():
                if opt("backup") and now - r["last_backup"] >= 86400:
                    try:
                        await run_backup(r["chat_id"], r["title"], r["fmt"], r["last_backup"], now)
                        db.execute("UPDATE backups SET last_backup=? WHERE chat_id=?", (now, r["chat_id"]))
                        db.commit()
                    except Exception:
                        log.exception("backup failed for %s", r["chat_id"])
            if now % 3600 < 60:
                purge_old()
        except Exception:
            log.exception("scheduler error")
        await asyncio.sleep(60)


FMT_MAP = {"html": "html", "xlsx": "xlsx", "excel": "xlsx", "both": "both"}


@client.on(events.NewMessage(outgoing=True, pattern=rf"(?s)^{P}backup(?:\s+(.*))?$"))
async def cmd_backup(event):
    args = (event.pattern_match.group(1) or "").split()
    sub = args[0].lower() if args else "help"
    fmt, target = "both", None
    for a in args[1:]:
        if a.lower() in FMT_MAP:
            fmt = FMT_MAP[a.lower()]
        else:
            target = a
    try:
        if sub == "list":
            rows = db.execute("SELECT * FROM backups").fetchall()
            if not rows:
                return await to_me(event, "📭 هیچ گروهی برای بکاپ ثبت نشده.")
            txt = "📦 گروه‌های بکاپ:\n" + "\n".join(
                f"• {r['title']} ({r['chat_id']}) — {r['fmt']} — آخرین بکاپ: "
                f"{fmt_date(r['last_backup'])} {fmt_time(r['last_backup'])}" for r in rows)
            return await to_me(event, txt)

        if sub not in ("add", "remove", "del", "now"):
            return await to_me(event, HELP)

        if target:
            ent = await client.get_entity(int(target) if re.fullmatch(r"-?\d+", target) else target)
        elif not event.is_private:
            ent = await event.get_chat()
        else:
            return await to_me(event, "⚠️ دستور را داخل گروه بزن یا آیدی/یوزرنیم گروه را بده.")
        cid = utils.get_peer_id(ent)
        title = utils.get_display_name(ent) or str(cid)
        now = int(time.time())

        if sub == "add":
            db.execute("INSERT OR REPLACE INTO backups VALUES (?,?,?,?)", (cid, title, fmt, now))
            db.commit()
            await to_me(event, f"✅ «{title}» به لیست بکاپ اضافه شد ({fmt}). اولین بکاپ ۲۴ ساعت دیگر.")
        elif sub in ("remove", "del"):
            db.execute("DELETE FROM backups WHERE chat_id=?", (cid,))
            db.commit()
            await to_me(event, f"🗑 «{title}» از لیست بکاپ حذف شد.")
        elif sub == "now":
            r = db.execute("SELECT * FROM backups WHERE chat_id=?", (cid,)).fetchone()
            since = r["last_backup"] if r else now - 86400
            use_fmt = r["fmt"] if r else fmt
            await to_me(event, f"⏳ در حال بکاپ‌گیری از «{title}» ...")
            n = await run_backup(cid, title, use_fmt, since, now)
            if r:
                db.execute("UPDATE backups SET last_backup=? WHERE chat_id=?", (now, cid))
                db.commit()
            await client.send_message("me", f"✅ بکاپ «{title}» ارسال شد ({n} پیام).")
    except Exception as ex:
        log.exception("cmd_backup")
        await to_me(event, f"❌ خطا: {ex}")


# --------------------------- متفرقه ---------------------------
HELP = f"""🤖 راهنمای سلف
{PREFIX}calc 2+2*3        ماشین‌حساب (sqrt, sin, cos, log, fact, pi, e)
{PREFIX}save              ریپلای روی مدیا → ذخیره در Saved Messages
{PREFIX}backup add [html|xlsx|both] [آیدی]
{PREFIX}backup remove [آیدی]
{PREFIX}backup list
{PREFIX}config                       نمایش تنظیمات ذخیره‌سازی (لیست گزینه‌ها)
{PREFIX}config all on|off            روشن/خاموش کردن همه (دستور کلی)
{PREFIX}config <گزینه|شماره> on|off|toggle   تنظیم تکی، مثلاً {PREFIX}config ttl off
{PREFIX}ai <سوال>                    پرسش از هوش مصنوعی پیش‌فرض (ریپلای هم جواب می‌دهد)
{PREFIX}claude / {PREFIX}gpt / {PREFIX}grok <سوال>   پرسش از مدل مشخص
{PREFIX}solve <مسئله/معادله>         حل مسئله‌ی ریاضی مرحله‌به‌مرحله (یا {PREFIX}حل)
{PREFIX}tr <زبان> <متن>              ترجمه و ارسال در همان چت (مثلاً {PREFIX}tr en سلام)
{PREFIX}tr <زبان>                    ریپلای روی پیام → ترجمه در Saved Messages
{PREFIX}aiset key <claude|gpt|grok> <کلید>   (فقط داخل Saved Messages)
{PREFIX}aiset model <claude|gpt|grok> <نام مدل>
{PREFIX}aiset default <claude|gpt|grok>
{PREFIX}aiset status
{PREFIX}pay <مبلغ تومان> [توضیح]     ساخت لینک پرداخت بله (در چت مشتری بزن)
{PREFIX}pay list                     آخرین درخواست‌های پرداخت
{PREFIX}bale token <توکن بات بله>    (فقط داخل Saved Messages)
{PREFIX}bale wallet <توکن کیف‌پول|test>  (فقط داخل Saved Messages)
{PREFIX}bale status                  وضعیت اتصال
{PREFIX}setlog [آیدی/یوزرنیم/me]     مقصد لاگ پیوی‌ها (داخل کانال بزن)
{PREFIX}setbackup [آیدی/یوزرنیم/me]  مقصد فایل‌های بکاپ
{PREFIX}settings                     نمایش تنظیمات فعلی
{PREFIX}backup now [آیدی]
{PREFIX}id                آیدی چت فعلی"""


@client.on(events.NewMessage(outgoing=True, pattern=rf"^{P}id$"))
async def cmd_id(event):
    await to_me(event, f"🆔 آیدی این چت: {event.chat_id}")


@client.on(events.NewMessage(outgoing=True, pattern=rf"^{P}help$"))
async def cmd_help(event):
    await to_me(event, HELP)


def get_setting(key, default=None):
    r = db.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    return r["value"] if r else default


def put_setting(key, value):
    db.execute("INSERT INTO settings(key,value) VALUES(?,?) "
               "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, str(value)))
    db.commit()


async def resolve_target(value):
    if value == "me":
        return await client.get_entity("me")
    return await client.get_entity(int(value) if re.fullmatch(r"-?\d+", value) else value)


@client.on(events.NewMessage(outgoing=True, pattern=rf"(?s)^{P}(setlog|setbackup)(?:\s+(\S+))?\s*$"))
async def cmd_set_target(event):
    """داخل کانال/گروه مقصد بزن، یا آیدی/یوزرنیم بده، یا 'me' برای Saved Messages."""
    global LOG_ENT, BACKUP_ENT
    which = event.pattern_match.group(1).lower()
    arg = event.pattern_match.group(2)
    try:
        if arg and arg.lower() in ("me", "saved"):
            ent = await client.get_entity("me")
        elif arg:
            ent = await client.get_entity(int(arg) if re.fullmatch(r"-?\d+", arg) else arg)
        elif not event.is_private:
            ent = await event.get_chat()
        else:
            return await to_me(event, f"⚠️ دستور {PREFIX}{which} را داخل کانال مقصد بزن "
                                      f"یا آیدی/یوزرنیم کانال را بده (یا me).")
        if getattr(ent, "id", None) == ME.id:
            value = "me"
        elif isinstance(ent, User):
            return await to_me(event, "⚠️ مقصد نمی‌تواند پیوی یک کاربر باشد؛ کانال یا گروه انتخاب کن.")
        else:
            value = str(utils.get_peer_id(ent))
        title = "Saved Messages" if value == "me" else (utils.get_display_name(ent) or value)
        if which == "setlog":
            put_setting("log_target", value)
            LOG_ENT = ent
            label = "لاگ پیوی‌ها"
        else:
            put_setting("backup_target", value)
            BACKUP_ENT = ent
            label = "بکاپ گروه‌ها"
        await to_me(event, f"✅ مقصد {label} روی «{title}» تنظیم شد.")
    except Exception as ex:
        log.exception("cmd_set_target")
        await to_me(event, f"❌ خطا: {ex}")


@client.on(events.NewMessage(outgoing=True, pattern=rf"^{P}settings$"))
async def cmd_settings(event):
    def name(e):
        return "Saved Messages" if e.id == ME.id else f"{utils.get_display_name(e)} ({utils.get_peer_id(e)})"
    await to_me(event, f"⚙️ تنظیمات فعلی\n📩 مقصد لاگ پیوی‌ها: {name(LOG_ENT)}\n"
                       f"📦 مقصد بکاپ: {name(BACKUP_ENT)}")


# ----------------------- ۵) درگاه پرداخت بله -----------------------
# جریان کار: .pay → لینک پرداخت در چت مشتری → مشتری در بله پرداخت می‌کند →
# بات بله (با getUpdates) رویداد پرداخت را می‌گیرد → به شما اطلاع می‌دهد.
BALE_API = "https://tapi.bale.ai"
BALE_TEST_WALLET = "WALLET-TEST-1111111111111111"
INVOICE_TTL = 24 * 3600  # اعتبار لینک پرداخت (ثانیه)

db.execute("""CREATE TABLE IF NOT EXISTS invoices(
  code TEXT PRIMARY KEY, title TEXT, description TEXT, amount INTEGER,
  status TEXT, tg_chat_id INTEGER, tg_msg_id INTEGER, created INTEGER,
  paid_at INTEGER, charge_id TEXT, payer TEXT)""")
db.commit()


async def bale_call(method, **params):
    token = get_setting("bale_token")
    if not token:
        raise RuntimeError("توکن بات بله تنظیم نشده است (.bale token ...)")
    try:
        async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=40)) as s:
            async with s.post(f"{BALE_API}/bot{token}/{method}", json=params) as r:
                data = await r.json(content_type=None)
    except (aiohttp.ClientError, asyncio.TimeoutError) as e:
        # عمداً آدرس (که توکن داخلش است) در پیام خطا نمی‌آید
        raise RuntimeError(f"{method}: خطای شبکه ({type(e).__name__})") from None
    if not isinstance(data, dict) or not data.get("ok"):
        desc = data.get("description") if isinstance(data, dict) else None
        raise RuntimeError(f"{method}: {desc or 'پاسخ نامعتبر'}")
    return data.get("result")


async def bale_username():
    u = get_setting("bale_username")
    if not u:
        me_ = await bale_call("getMe")
        u = (me_ or {}).get("username")
        if not u:
            raise RuntimeError("بات بله یوزرنیم ندارد")
        put_setting("bale_username", u)
    return u


def get_invoice(code):
    return db.execute("SELECT * FROM invoices WHERE code=?", (code or "",)).fetchone()


def fmt_payer(user):
    user = user or {}
    name = " ".join(x for x in (user.get("first_name"), user.get("last_name")) if x) or "نامشخص"
    uname = f" @{user['username']}" if user.get("username") else ""
    return f"{name}{uname} (id: {user.get('id', '؟')})"


@client.on(events.NewMessage(outgoing=True, pattern=rf"(?s)^{P}pay(?:\s+(.*))?$"))
async def cmd_pay(event):
    args = (event.pattern_match.group(1) or "").strip()
    if args.lower() == "list":
        rows = db.execute("SELECT * FROM invoices ORDER BY created DESC LIMIT 10").fetchall()
        if not rows:
            return await to_me(event, "📭 هنوز درخواست پرداختی ساخته نشده.")
        icon = {"pending": "⏳", "paid": "✅", "mismatch": "⚠️"}
        txt = "💳 آخرین درخواست‌های پرداخت:\n" + "\n".join(
            f"{icon.get(r['status'], '•')} {r['amount'] // 10:,} تومان — {r['description']} — "
            f"{fmt_date(r['created'])} {fmt_time(r['created'])}" for r in rows)
        return await to_me(event, txt)

    m = re.match(r"^([\d,٬]+)\s*(.*)$", args.translate(DIGITS), re.S)
    if not m:
        return await to_me(event, f"⚠️ فرمت: {PREFIX}pay <مبلغ به تومان> [توضیحات]\n"
                                  f"مثال: {PREFIX}pay 50000 خرید اشتراک")
    toman = int(re.sub(r"[,٬]", "", m.group(1)))
    desc = (m.group(2) or "").strip() or "پرداخت"
    if toman <= 0:
        return await to_me(event, "⚠️ مبلغ باید بیشتر از صفر باشد.")
    if not get_setting("bale_token") or not get_setting("bale_wallet"):
        return await to_me(event, f"⚠️ اول توکن بات و کیف‌پول را تنظیم کن: "
                                  f"{PREFIX}bale token ... و {PREFIX}bale wallet ... (داخل Saved Messages)")

    code = secrets.token_urlsafe(9)
    title, description, rials = desc[:32], desc[:255], toman * 10  # بله مبلغ را به ریال می‌گیرد
    db.execute("INSERT INTO invoices(code,title,description,amount,status,tg_chat_id,tg_msg_id,created) "
               "VALUES (?,?,?,?,?,?,?,?)",
               (code, title, description, rials, "pending", event.chat_id, event.id, int(time.time())))
    db.commit()
    try:
        link = None
        try:  # لینک مستقیم فاکتور؛ اگر پارامترها قبول نشد به لینک بات برمی‌گردیم
            res = await bale_call("createInvoiceLink", title=title, description=description,
                                  payload=code, provider_token=get_setting("bale_wallet"),
                                  prices=[{"label": title, "amount": rials}])
            if isinstance(res, str) and res.startswith("http"):
                link = res
        except Exception as ex:
            log.info("createInvoiceLink unavailable, using bot deep link: %s", ex)
        if not link:
            link = f"https://ble.ir/{await bale_username()}?start={code}"
    except Exception as ex:
        db.execute("DELETE FROM invoices WHERE code=?", (code,))
        db.commit()
        return await to_me(event, f"❌ ساخت لینک پرداخت ناموفق بود: {ex}")

    SELF_EDITED.add(event.id)
    await event.edit(f"💳 درخواست پرداخت\n📝 {desc}\n💰 مبلغ: {toman:,} تومان\n\n"
                     f"🔗 پرداخت با بله:\n{link}\n\n⏳ اعتبار لینک: ۲۴ ساعت")


@client.on(events.NewMessage(outgoing=True, pattern=rf"(?s)^{P}bale(?:\s+(.*))?$"))
async def cmd_bale(event):
    args = (event.pattern_match.group(1) or "").split()
    sub = args[0].lower() if args else "status"
    try:
        if sub in ("token", "wallet"):
            if event.chat_id != ME.id:  # توکن نباید داخل چت دیگران نوشته شود
                return await to_me(event, "⚠️ این دستور را فقط داخل Saved Messages بزن. "
                                          "پیام پاک شد، ولی اگر توکن را در چت دیگری نوشتی "
                                          "فوراً از BotFather بله توکن را عوض کن.")
            if len(args) < 2:
                return await to_me(event, f"⚠️ فرمت: {PREFIX}bale {sub} <مقدار>")
            val = args[1]
            if sub == "token":
                old = get_setting("bale_token")
                put_setting("bale_token", val)
                try:
                    me_ = await bale_call("getMe")
                except Exception as ex:
                    if old:
                        put_setting("bale_token", old)
                    else:
                        db.execute("DELETE FROM settings WHERE key='bale_token'")
                        db.commit()
                    return await to_me(event, f"❌ توکن معتبر نیست: {ex}")
                put_setting("bale_username", (me_ or {}).get("username") or "")
                put_setting("bale_offset", 0)
                return await to_me(event, f"✅ بات بله متصل شد: @{(me_ or {}).get('username', '?')}")
            test = val.lower() == "test"
            put_setting("bale_wallet", BALE_TEST_WALLET if test else val)
            return await to_me(event, "✅ توکن کیف‌پول ذخیره شد" + (" (حالت تست)." if test else "."))

        tok, wal = get_setting("bale_token"), get_setting("bale_wallet")
        un = get_setting("bale_username")
        await to_me(event,
                    "💳 وضعیت درگاه بله\n"
                    f"🤖 توکن بات: {'✅ @' + un if tok and un else ('✅' if tok else '❌')}\n"
                    f"👛 توکن کیف‌پول: "
                    f"{'❌' if not wal else ('✅ (تست)' if wal == BALE_TEST_WALLET else '✅')}")
    except Exception as ex:
        log.exception("cmd_bale")
        await to_me(event, f"❌ خطا: {ex}")


async def handle_bale_update(u):
    wallet = get_setting("bale_wallet")

    q = u.get("pre_checkout_query")
    if q:
        inv = get_invoice(q.get("invoice_payload"))
        amount = q.get("total_amount")
        ok = bool(inv and inv["status"] == "pending"
                  and time.time() - inv["created"] <= INVOICE_TTL
                  and (amount is None or amount == inv["amount"]))
        if ok:
            await bale_call("answerPreCheckoutQuery", pre_checkout_query_id=q["id"], ok=True)
        else:
            await bale_call("answerPreCheckoutQuery", pre_checkout_query_id=q["id"], ok=False,
                            error_message="این درخواست پرداخت معتبر نیست یا منقضی شده است.")
        return

    m = u.get("message")
    if not m:
        return
    chat_id = (m.get("chat") or {}).get("id")

    sp = m.get("successful_payment")
    if sp:
        inv = get_invoice(sp.get("invoice_payload"))
        if not inv or inv["status"] == "paid":
            return
        amount = sp.get("total_amount")
        if amount is not None and amount != inv["amount"]:
            db.execute("UPDATE invoices SET status='mismatch' WHERE code=?", (inv["code"],))
            db.commit()
            return await client.send_message(
                "me", f"⚠️ پرداخت با مبلغ نامطابق دریافت شد!\nانتظار: {inv['amount']} ریال | "
                      f"دریافتی: {amount}\n📝 {inv['description']}\nبررسی دستی لازم است.")
        payer = fmt_payer(m.get("from"))
        charge = str(sp.get("provider_payment_charge_id") or sp.get("telegram_payment_charge_id") or "")
        now = int(time.time())
        db.execute("UPDATE invoices SET status='paid', paid_at=?, charge_id=?, payer=? WHERE code=?",
                   (now, charge, payer, inv["code"]))
        db.commit()
        toman = inv["amount"] // 10
        await client.send_message(
            "me", f"✅ پرداخت موفق\n💰 {toman:,} تومان\n📝 {inv['description']}\n"
                  f"👤 پرداخت‌کننده (بله): {payer}\n🧾 کد رهگیری: {charge or '—'}\n"
                  f"🕒 {fmt_time(now)} - {fmt_date(now)}")
        try:  # پیام لینک پرداخت در چت مشتری به «پرداخت شد» تغییر می‌کند
            SELF_EDITED.add(inv["tg_msg_id"])
            await client.edit_message(inv["tg_chat_id"], inv["tg_msg_id"],
                                      f"✅ پرداخت شد\n📝 {inv['description']}\n💰 مبلغ: {toman:,} تومان")
        except Exception:
            log.warning("could not edit payment message", exc_info=True)
        try:
            await bale_call("sendMessage", chat_id=chat_id, text="✅ پرداخت شما با موفقیت ثبت شد. ممنون از شما!")
        except Exception:
            pass
        return

    text = (m.get("text") or "").strip()
    if text.startswith("/start"):
        parts = text.split(maxsplit=1)
        inv = get_invoice(parts[1].strip() if len(parts) > 1 else "")
        if not inv or inv["status"] != "pending" or time.time() - inv["created"] > INVOICE_TTL:
            return await bale_call("sendMessage", chat_id=chat_id,
                                   text="این لینک پرداخت معتبر نیست یا منقضی شده است.")
        await bale_call("sendInvoice", chat_id=chat_id, title=inv["title"],
                        description=inv["description"], payload=inv["code"],
                        provider_token=wallet,
                        prices=[{"label": inv["title"], "amount": inv["amount"]}])


async def bale_poller():
    offset = int(get_setting("bale_offset", "0") or 0)
    webhook_cleared = False
    while True:
        if not get_setting("bale_token") or not get_setting("bale_wallet"):
            await asyncio.sleep(10)
            continue
        try:
            if not webhook_cleared:
                try:
                    await bale_call("deleteWebhook")
                except Exception:
                    pass
                webhook_cleared = True
            updates = await bale_call("getUpdates", offset=offset, limit=50) or []
            for u in updates:
                offset = u["update_id"] + 1
                put_setting("bale_offset", offset)
                try:
                    await handle_bale_update(u)
                except Exception:
                    log.exception("bale update failed")
            await asyncio.sleep(0 if updates else 2)
        except Exception as ex:
            log.warning("bale poll error: %s", ex)
            await asyncio.sleep(10)


# --------------- ۶) هوش مصنوعی: پرسش‌وپاسخ، ریاضی، ترجمه ---------------
# کلیدهای API با دستور .aiset در دیتابیس ذخیره می‌شوند (نه داخل سورس).
AI_PROVIDERS = {
    "claude": {"label": "Claude", "model": "claude-sonnet-5-5"},
    "gpt": {"label": "ChatGPT", "model": "chat-latest"},
    "grok": {"label": "Grok", "model": "grok-4.7"},
}
AI_ALIASES = {"claude": "claude", "anthropic": "claude",
              "gpt": "gpt", "chatgpt": "gpt", "openai": "gpt",
              "grok": "grok", "xai": "grok"}

SYS_GENERAL = ("You are a helpful, accurate assistant. Answer in the same language as the "
               "user's question. Be clear and reasonably concise. Telegram does not render "
               "markdown or LaTeX, so use plain text.")
SYS_MATH = ("You are an expert mathematician. Solve the user's problem step by step, double-check "
            "every calculation, and state the final answer clearly at the end. Cover equations, "
            "systems, inequalities, radicals, algebra, calculus, geometry, probability and word "
            "problems. Reply in the user's language (Persian if unsure). Telegram does not render "
            "LaTeX, so write math in plain text, e.g. x^2, sqrt(8), 3/4, ∫, Σ.")


def sys_translate(lang):
    return (f"You are a professional translator. Translate the user's text into {lang}. "
            "Detect the source language automatically. Be accurate and natural: preserve meaning, "
            "tone, names, numbers, emojis, formatting and line breaks. Output ONLY the "
            "translation, with no explanations or quotation marks.")


async def ai_ask(provider, prompt, system):
    label = AI_PROVIDERS[provider]["label"]
    key = get_setting(f"ai_key_{provider}")
    if not key:
        raise RuntimeError(f"کلید API برای {label} تنظیم نشده ({PREFIX}aiset key {provider} ...)")
    model = get_setting(f"ai_model_{provider}") or AI_PROVIDERS[provider]["model"]
    try:
        async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=180)) as s:
            if provider == "claude":
                url = "https://api.anthropic.com/v1/messages"
                headers = {"x-api-key": key, "anthropic-version": "2023-06-01"}
                body = {"model": model, "max_tokens": 4096, "system": system,
                        "messages": [{"role": "user", "content": prompt}]}
            else:
                base = "https://api.openai.com/v1" if provider == "gpt" else "https://api.x.ai/v1"
                url = f"{base}/chat/completions"
                headers = {"Authorization": f"Bearer {key}"}
                body = {"model": model,
                        "messages": [{"role": "system", "content": system},
                                     {"role": "user", "content": prompt}]}
            async with s.post(url, json=body, headers=headers) as r:
                status = r.status
                data = await r.json(content_type=None)
    except (aiohttp.ClientError, asyncio.TimeoutError) as e:
        raise RuntimeError(f"{label}: خطای شبکه ({type(e).__name__})") from None
    if status >= 400 or not isinstance(data, dict):
        err = data.get("error") if isinstance(data, dict) else None
        msg = err.get("message") if isinstance(err, dict) else (err or f"HTTP {status}")
        raise RuntimeError(f"{label}: {msg}")
    if provider == "claude":
        text = "".join(b.get("text", "") for b in data.get("content", []) if b.get("type") == "text")
    else:
        text = ((data.get("choices") or [{}])[0].get("message") or {}).get("content") or ""
    if not text.strip():
        raise RuntimeError(f"{label}: پاسخ خالی دریافت شد")
    return text.strip()


def default_provider():
    p = get_setting("ai_default")
    if p in AI_PROVIDERS:
        return p
    for name in AI_PROVIDERS:
        if get_setting(f"ai_key_{name}"):
            return name
    return "claude"


async def to_me_long(event, text):
    try:
        await event.delete()
    except Exception:
        pass
    for i in range(0, len(text), 4000):
        await client.send_message("me", text[i:i + 4000])


@client.on(events.NewMessage(outgoing=True, pattern=rf"(?s)^{P}(ai|claude|gpt|grok|solve|حل)(?:\s+(.+))?$"))
async def cmd_ai(event):
    cmd = event.pattern_match.group(1).lower()
    q = (event.pattern_match.group(2) or "").strip()
    if not q:  # روی پیام دیگری ریپلای کرده؟ همان متن سوال است
        r = await event.get_reply_message()
        if r and r.raw_text:
            q = r.raw_text.strip()
    if not q:
        return await to_me(event, f"⚠️ سوالت را بعد از دستور بنویس یا روی یک پیام ریپلای کن.\nمثال: {PREFIX}{cmd} ...")
    solve = cmd in ("solve", "حل")
    prov = cmd if cmd in AI_PROVIDERS else default_provider()
    try:
        ans = await ai_ask(prov, q, SYS_MATH if solve else SYS_GENERAL)
    except Exception as ex:
        return await to_me(event, f"❌ {ex}")
    await to_me_long(event, f"{'🧮' if solve else '❓'} {q}\n\n🤖 {AI_PROVIDERS[prov]['label']}:\n{ans}")


@client.on(events.NewMessage(outgoing=True, pattern=rf"(?s)^{P}tr(?:\s+(\S+)(?:\s+(.+))?)?$"))
async def cmd_tr(event):
    lang = event.pattern_match.group(1)
    text = (event.pattern_match.group(2) or "").strip()
    if not lang:
        return await to_me(event, f"⚠️ فرمت: {PREFIX}tr <زبان مقصد> <متن>\nمثال: {PREFIX}tr en سلام خوبی؟\n"
                                  f"یا روی یک پیام ریپلای کن و بزن: {PREFIX}tr fa")
    reply_mode = False
    if not text:
        r = await event.get_reply_message()
        if r and r.raw_text:
            text, reply_mode = r.raw_text.strip(), True
    if not text:
        return await to_me(event, "⚠️ متنی برای ترجمه پیدا نشد.")
    try:
        ans = await ai_ask(default_provider(), text, sys_translate(lang))
    except Exception as ex:
        return await to_me(event, f"❌ {ex}")
    if reply_mode or event.chat_id == ME.id:
        return await to_me_long(event, f"🌐 ترجمه به {lang}\n\n📄 اصل:\n{text}\n\n✅ ترجمه:\n{ans}")
    # ترجمه‌ی متن خودت: پیام دستور پاک می‌شود و ترجمه در همان چت فرستاده می‌شود
    try:
        await event.delete()
    except Exception:
        pass
    for i in range(0, len(ans), 4000):
        await client.send_message(event.chat_id, ans[i:i + 4000])


@client.on(events.NewMessage(outgoing=True, pattern=rf"(?s)^{P}aiset(?:\s+(.*))?$"))
async def cmd_aiset(event):
    args = (event.pattern_match.group(1) or "").split()
    sub = args[0].lower() if args else "status"
    try:
        if sub in ("key", "model", "default"):
            if sub == "default":
                if len(args) < 2 or args[1].lower() not in AI_ALIASES:
                    return await to_me(event, f"⚠️ فرمت: {PREFIX}aiset default <claude|gpt|grok>")
                prov = AI_ALIASES[args[1].lower()]
                put_setting("ai_default", prov)
                return await to_me(event, f"✅ هوش مصنوعی پیش‌فرض: {AI_PROVIDERS[prov]['label']}")
            if len(args) < 3 or args[1].lower() not in AI_ALIASES:
                return await to_me(event, f"⚠️ فرمت: {PREFIX}aiset {sub} <claude|gpt|grok> <مقدار>")
            prov = AI_ALIASES[args[1].lower()]
            label = AI_PROVIDERS[prov]["label"]
            if sub == "key":
                if event.chat_id != ME.id:  # کلید نباید داخل چت دیگران نوشته شود
                    return await to_me(event, "⚠️ این دستور را فقط داخل Saved Messages بزن. پیام پاک شد، "
                                              "ولی اگر کلید را جای دیگری نوشتی فوراً کلید جدید بساز.")
                put_setting(f"ai_key_{prov}", args[2])
                if not get_setting("ai_default"):
                    put_setting("ai_default", prov)
                return await to_me(event, f"✅ کلید {label} ذخیره شد.")
            put_setting(f"ai_model_{prov}", args[2])
            return await to_me(event, f"✅ مدل {label}: {args[2]}")

        dflt = default_provider()
        lines = ["🤖 تنظیمات هوش مصنوعی"]
        for name, info in AI_PROVIDERS.items():
            model = get_setting(f"ai_model_{name}") or info["model"]
            ok = "✅" if get_setting(f"ai_key_{name}") else "❌"
            star = " ⭐ پیش‌فرض" if name == dflt else ""
            lines.append(f"{ok} {info['label']} — مدل: {model}{star}")
        await to_me(event, "\n".join(lines))
    except Exception as ex:
        log.exception("cmd_aiset")
        await to_me(event, f"❌ خطا: {ex}")


# ----------------- ۷) تنظیم این‌که چه چیزهایی ذخیره شوند -----------------
# مقدار دوم هر گزینه پیش‌فرض است؛ با دستور .config در دیتابیس تغییر می‌کند.
OPTIONS = {
    "incoming": ("📩 لاگ پیام‌های دریافتی پیوی", True),
    "outgoing": ("📤 لاگ پیام‌های ارسالی خودت در پیوی", LOG_OUTGOING),
    "edits": ("✏️ لاگ ویرایش پیام‌ها", True),
    "deletes": ("🗑 لاگ پیام‌های پاک‌شده", True),
    "media": ("🖼 ذخیره و ارسال مدیا در لاگ", CACHE_MEDIA),
    "ttl": ("💣 ذخیره‌ی مدیای نابودشونده", True),
    "bots": ("🤖 لاگ پیوی ربات‌ها", not IGNORE_BOTS),
    "backup": ("📦 بکاپ خودکار گروه‌ها", True),
}
ON_WORDS = {"on", "1", "yes", "true", "روشن", "فعال"}
OFF_WORDS = {"off", "0", "no", "false", "خاموش", "غیرفعال"}
TOGGLE_WORDS = {"toggle", "t", "تغییر"}


def opt(name):
    v = get_setting(f"opt_{name}")
    if v is None:
        return bool(OPTIONS[name][1])
    return v == "1"


def options_text():
    lines = ["⚙️ تنظیمات ذخیره‌سازی"]
    for i, (n, (label, _)) in enumerate(OPTIONS.items(), 1):
        lines.append(f"{i}. {'✅' if opt(n) else '❌'} {n} — {label}")
    lines.append(f"\n{PREFIX}config all off  |  {PREFIX}config ttl on  |  {PREFIX}config 3 toggle")
    return "\n".join(lines)


@client.on(events.NewMessage(outgoing=True, pattern=rf"(?s)^{P}config(?:\s+(.*))?$"))
async def cmd_config(event):
    args = (event.pattern_match.group(1) or "").split()
    names = list(OPTIONS)
    if not args or args[0].lower() == "status":
        return await to_me(event, options_text())
    key = args[0].lower()
    if key.isdigit() and 1 <= int(key) <= len(names):
        key = names[int(key) - 1]
    word = args[1].lower() if len(args) > 1 else "toggle"

    if key == "all":  # دستور کلی: همه‌ی گزینه‌ها
        if word in ON_WORDS:
            state = "1"
        elif word in OFF_WORDS:
            state = "0"
        else:
            return await to_me(event, f"⚠️ فرمت: {PREFIX}config all on  یا  {PREFIX}config all off")
        for n in names:
            put_setting(f"opt_{n}", state)
        return await to_me(event, ("✅ همه روشن شد.\n\n" if state == "1" else "⛔ همه خاموش شد.\n\n") + options_text())

    if key not in OPTIONS:  # دستور تکی
        return await to_me(event, f"⚠️ گزینه‌ی «{args[0]}» وجود ندارد.\n\n" + options_text())
    if word in ON_WORDS:
        new = True
    elif word in OFF_WORDS:
        new = False
    elif word in TOGGLE_WORDS:
        new = not opt(key)
    else:
        return await to_me(event, f"⚠️ فرمت: {PREFIX}config {key} on|off|toggle")
    put_setting(f"opt_{key}", "1" if new else "0")
    await to_me(event, f"{'✅ روشن' if new else '⛔ خاموش'} شد: {OPTIONS[key][0]}\n\n" + options_text())


async def setup():
    global ME, LOG_ENT, BACKUP_ENT
    ME = await client.get_me()
    await client.get_dialogs()
    try:
        LOG_ENT = await resolve_target(get_setting("log_target", str(LOG_CHANNEL)))
    except Exception:
        log.warning("log target not reachable, falling back to Saved Messages", exc_info=True)
        LOG_ENT = await client.get_entity("me")
    try:
        BACKUP_ENT = await resolve_target(get_setting("backup_target", str(BACKUP_TARGET)))
    except Exception:
        log.warning("backup target not reachable, falling back to Saved Messages", exc_info=True)
        BACKUP_ENT = await client.get_entity("me")
    purge_old()
    client.loop.create_task(scheduler())
    client.loop.create_task(bale_poller())
    log.info("Self is running as %s (%s)", utils.get_display_name(ME), ME.id)


if __name__ == "__main__":
    client.start()  # بار اول شماره و کد تأیید را می‌پرسد
    client.loop.run_until_complete(setup())
    client.run_until_disconnected()
